#!/usr/bin/env python3
"""Remove duplicate rows from the scraped bill files.

    python pipeline/dedupe_bills.py           every live session's bill file ([[bills.sessions]] details in config.toml)

The scraper sometimes saves a bill twice: an early snapshot (often just "Filed with the Clerk") and its current row.
For each bill that appears more than once, this keeps the row with the latest action date (then the longest action
history) and deletes the others. If only a deleted row has a code, the code moves to the kept row first; if both have
different codes, the kept row's code stays (the one the site already uses) and the other is listed.

Only the duplicate rows change: the file is edited in place with openpyxl, so other rows, columns and formatting are
left as they are. On GitHub, run "Update website" with "Remove duplicate bills" ticked; it saves the file back to Dropbox.
"""
import os, re, sys, tomllib
from collections import defaultdict
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
DATA = Path(os.environ.get('IL_AGENDAS_DATA') or ROOT).expanduser().resolve()


def dedupe(path, col):
    """Edits the workbook at path; returns (rows removed, log lines)."""
    wb = openpyxl.load_workbook(path)
    ws = wb.worksheets[0]
    head = {c.value: i + 1 for i, c in enumerate(ws[1]) if c.value is not None}
    need = [col['bill'], col['last_action_date'], col['actions'], col['subtopic']]
    missing = [c for c in need if c not in head]
    if missing:
        raise SystemExit(f'{path.name}: missing columns {", ".join(missing)}')
    major = head.get('major_code')   # moved along with the subtopic code when a deleted row has the only code
    val = lambda r, c: ws.cell(r, head[c]).value
    rows = defaultdict(list)
    for r in range(2, ws.max_row + 1):
        b = val(r, col['bill'])
        if b is not None and str(b).strip():
            rows[re.sub(r'\s+', ' ', str(b)).strip()].append(r)

    def newest(r):
        d = val(r, col['last_action_date'])
        d = d.timestamp() / 86400 if hasattr(d, 'timestamp') else float(d or 0) if str(d or '').replace('.', '', 1).isdigit() else 0
        return (d, len(str(val(r, col['actions']) or '')))

    drop, log = [], []
    empty = lambda v: v is None or str(v).strip() == ''
    for b, rs in sorted(rows.items()):
        if len(rs) < 2:
            continue
        keep = max(rs, key=newest)
        for r in rs:
            if r == keep:
                continue
            code, kept = val(r, col['subtopic']), val(keep, col['subtopic'])
            note = ''
            if empty(kept) and not empty(code):
                ws.cell(keep, head[col['subtopic']]).value = code
                if major:
                    ws.cell(keep, major).value = ws.cell(r, major).value
                note = f'; moved its code {code} to the kept row'
            elif not empty(code) and code != kept:
                note = f'; its code {code} differs, kept {kept}'
            drop.append(r)
            log.append(f'{b}: removed row {r}, kept row {keep}{note}')
    for r in sorted(drop, reverse=True):
        ws.delete_rows(r)
    if drop:
        wb.save(path)
    return len(drop), log


def main():
    cfg = tomllib.loads((ROOT / 'config.toml').read_text(encoding='utf-8'))
    col = cfg['bills']['columns']
    for s in cfg['bills']['sessions']:
        path = DATA / s['details']
        if not path.exists():
            continue
        n, log = dedupe(path, col)
        print(f'{path.name}: ' + (f'removed {n} duplicate row{"s" if n != 1 else ""}.' if n else 'no duplicates.'))
        for line in log:
            print('  ' + line)


if __name__ == '__main__':
    sys.exit(main())
